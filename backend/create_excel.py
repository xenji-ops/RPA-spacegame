"""
create_excel.py
Run this ONCE to generate the planet_data.xlsx file that the RPA bot reads.
The Excel file is the "database" the UiPath / Python RPA bot processes.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def create_planet_database():
    wb = Workbook()
    ws = wb.active
    ws.title = "PlanetDatabase"

    # ── Column widths ──────────────────────────────────────────────────────────
    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 28
    ws.column_dimensions['D'].width = 14
    ws.column_dimensions['E'].width = 14
    ws.column_dimensions['F'].width = 14
    ws.column_dimensions['G'].width = 20

    # ── Styles ─────────────────────────────────────────────────────────────────
    header_font  = Font(bold=True, color="FFFFFF", size=11)
    title_font   = Font(bold=True, color="00F5FF", size=14)
    danger_fill  = PatternFill("solid", fgColor="4A0000")
    safe_fill    = PatternFill("solid", fgColor="004A22")
    header_fill  = PatternFill("solid", fgColor="050D2E")
    alt_fill     = PatternFill("solid", fgColor="0A1A40")
    thin_border  = Border(
        left=Side(style='thin', color='334466'),
        right=Side(style='thin', color='334466'),
        top=Side(style='thin', color='334466'),
        bottom=Side(style='thin', color='334466'),
    )

    center = Alignment(horizontal='center', vertical='center')
    left   = Alignment(horizontal='left',   vertical='center')

    # ── Title row ──────────────────────────────────────────────────────────────
    ws.merge_cells('A1:G1')
    ws['A1'] = '🌌  GALAXY RESCUE — PLANETARY DATABASE  🌌'
    ws['A1'].font      = title_font
    ws['A1'].fill      = header_fill
    ws['A1'].alignment = center
    ws.row_dimensions[1].height = 28

    # ── Thresholds row ────────────────────────────────────────────────────────
    ws.merge_cells('A2:G2')
    ws['A2'] = 'DANGER THRESHOLDS:  Oxygen < 35%  |  Energy < 25%  |  Shield < 30%'
    ws['A2'].font      = Font(bold=True, color="FF3366", size=10)
    ws['A2'].fill      = PatternFill("solid", fgColor="1A0010")
    ws['A2'].alignment = center
    ws.row_dimensions[2].height = 20

    # ── Header row ────────────────────────────────────────────────────────────
    headers = ['ID', 'Planet Name', 'Description', 'Oxygen (%)', 'Energy (%)', 'Shield (%)', 'Status']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = center
        cell.border    = thin_border
    ws.row_dimensions[3].height = 22

    # ── Planet data rows ──────────────────────────────────────────────────────
    # This matches EXACTLY the data in gameData.js so the bot results are consistent
    planets = [
        ('P01', 'Zephyra',  'Ocean world with dense atmosphere',    28, 72, 85),
        ('P02', 'Volcanis', 'Volcanic gas giant, unstable core',    61, 18, 22),
        ('P03', 'Crystara', 'Crystal nebula, stable conditions',    78, 64, 71),
    ]

    # Thresholds (must match gameData.js DANGER_THRESHOLDS)
    O2_THRESH   = 35
    ENG_THRESH  = 25
    SHD_THRESH  = 30

    for row_idx, (pid, name, desc, oxy, eng, shd) in enumerate(planets, 4):
        in_danger = oxy < O2_THRESH or eng < ENG_THRESH or shd < SHD_THRESH
        status    = '🔴 IN DANGER' if in_danger else '🟢 SAFE'
        row_fill  = danger_fill if in_danger else safe_fill

        data = [pid, name, desc, oxy, eng, shd, status]
        for col, val in enumerate(data, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.fill   = row_fill
            cell.border = thin_border
            cell.alignment = center if col != 3 else left
            # Highlight critical values in red
            if col == 4 and oxy  < O2_THRESH:  cell.font = Font(bold=True, color="FF3366")
            elif col == 5 and eng < ENG_THRESH: cell.font = Font(bold=True, color="FF3366")
            elif col == 6 and shd < SHD_THRESH: cell.font = Font(bold=True, color="FF3366")
            else: cell.font = Font(color="E8F4FF")
        ws.row_dimensions[row_idx].height = 22

    # ── Second sheet: Thresholds reference ───────────────────────────────────
    ws2 = wb.create_sheet("Thresholds")
    ws2['A1'] = 'Parameter'
    ws2['B1'] = 'Critical Below (%)'
    ws2['C1'] = 'Notes'
    ws2['A2'] = 'Oxygen'  ; ws2['B2'] = O2_THRESH  ; ws2['C2'] = 'Breathability minimum'
    ws2['A3'] = 'Energy'  ; ws2['B3'] = ENG_THRESH  ; ws2['C3'] = 'Core stability minimum'
    ws2['A4'] = 'Shield'  ; ws2['B4'] = SHD_THRESH  ; ws2['C4'] = 'Deflector integrity minimum'
    for row in ws2.iter_rows(min_row=1, max_row=4, min_col=1, max_col=3):
        for cell in row:
            cell.font   = Font(color="E8F4FF", bold=(cell.row == 1))
            cell.fill   = header_fill
            cell.border = thin_border
            cell.alignment = center

    output_path = 'planet_data.xlsx'
    wb.save(output_path)
    print(f'[OK] Created {output_path}')
    print('   Sheet 1: PlanetDatabase -- planet readings the RPA bot will scan')
    print('   Sheet 2: Thresholds    -- danger threshold reference')
    return output_path

if __name__ == '__main__':
    create_planet_database()
